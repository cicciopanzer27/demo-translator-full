from docx import Document
from typing import List, Dict
import logging

logger = logging.getLogger(__name__)

class DOCXExtractor:
    """Estrae testo da DOCX preservando struttura"""
    
    def extract_with_structure(self, docx_path: str, progress_callback=None) -> List[Dict]:
        """
        Estrae testo preservando paragrafi e stili
        
        Returns:
            List[Dict]: [{
                'type': 'paragraph' | 'heading' | 'table',
                'text': str,
                'style': str,
                'level': int  # per heading
            }]
        """
        doc = Document(docx_path)
        total_elements = len(doc.paragraphs) + len(doc.tables)
        elements = []
        processed = 0
        
        # Paragrafi e heading
        for para in doc.paragraphs:
            processed += 1
            if progress_callback:
                progress_callback(processed, total_elements, "Estrazione contenuto")
            
            if not para.text.strip():
                continue
            
            element = {
                'type': 'heading' if para.style.name.startswith('Heading') else 'paragraph',
                'text': para.text,
                'style': para.style.name,
                'bold': para.runs[0].bold if para.runs else False,
                'italic': para.runs[0].italic if para.runs else False
            }
            
            if element['type'] == 'heading':
                element['level'] = int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1
            
            elements.append(element)
        
        # Tabelle
        for table in doc.tables:
            processed += 1
            if progress_callback:
                progress_callback(processed, total_elements, "Estrazione tabelle")
            
            table_text = []
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                table_text.append(row_text)
            
            elements.append({
                'type': 'table',
                'text': str(table_text),  # Serializzato
                'rows': len(table.rows),
                'cols': len(table.columns)
            })
        
        logger.info(f"Estratti {len(elements)} elementi")
        return elements
    
    def extract_simple(self, docx_path: str) -> str:
        """Estrazione semplice tutto il testo"""
        doc = Document(docx_path)
        return "\n".join([para.text for para in doc.paragraphs])

