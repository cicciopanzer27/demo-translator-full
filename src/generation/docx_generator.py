from docx import Document
from docx.shared import Pt, Inches
from typing import List, Dict
import logging

logger = logging.getLogger(__name__)

class DOCXGenerator:
    """Genera DOCX da testo tradotto con preservazione struttura"""
    
    def create_docx(self, pages_data: List[Dict], output_path: str, progress_callback=None):
        """
        Crea DOCX da pagine tradotte preservando struttura quando disponibile
        """
        # Verifica se abbiamo struttura DOCX preservata
        has_structure = any('elements' in page for page in pages_data)
        
        if has_structure:
            self._create_docx_with_structure(pages_data, output_path, progress_callback)
        else:
            self._create_docx_simple(pages_data, output_path, progress_callback)
    
    def _create_docx_simple(self, pages_data: List[Dict], output_path: str, progress_callback=None):
        """Creazione DOCX standard"""
        doc = Document()
        
        # Setup margini
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        total_pages = len(pages_data)
        
        for i, page_data in enumerate(pages_data):
            if progress_callback:
                progress_callback(
                    (i / total_pages) * 100,
                    f"Generazione pagina {i+1}/{total_pages}"
                )
            
            text = page_data['text']
            
            # Split in paragrafi
            paragraphs = text.split('\n\n')
            
            for para_text in paragraphs:
                if para_text.strip():
                    p = doc.add_paragraph(para_text)
                    
                    # Formattazione
                    p.style = 'Normal'
                    for run in p.runs:
                        run.font.size = Pt(11)
            
            # Page break (tranne ultima)
            if i < total_pages - 1:
                doc.add_page_break()
        
        if progress_callback:
            progress_callback(90, "Salvataggio DOCX...")
        
        doc.save(output_path)
        
        logger.info(f"DOCX generato (modalità standard): {output_path}")
    
    def _create_docx_with_structure(self, pages_data: List[Dict], output_path: str, progress_callback=None):
        """Creazione DOCX preservando struttura originale (heading, style, etc)"""
        doc = Document()
        
        # Setup margini
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        total_pages = len(pages_data)
        
        for i, page_data in enumerate(pages_data):
            if progress_callback:
                progress_callback(
                    (i / total_pages) * 100,
                    f"Generazione pagina {i+1}/{total_pages} (con struttura)"
                )
            
            # Se abbiamo elementi strutturati, usali
            if 'elements' in page_data:
                elements = page_data['elements']
                
                for element in elements:
                    elem_type = element.get('type', 'paragraph')
                    text = element.get('text', '').strip()
                    
                    if not text:
                        continue
                    
                    # Crea paragrafo appropriato
                    if elem_type == 'heading':
                        level = element.get('level', 1)
                        p = doc.add_heading(text, level=level)
                    else:
                        p = doc.add_paragraph(text)
                        
                        # Applica formattazione preservata
                        if element.get('bold', False) and p.runs:
                            p.runs[0].bold = True
                        if element.get('italic', False) and p.runs:
                            p.runs[0].italic = True
                        
                        # Font size
                        for run in p.runs:
                            run.font.size = Pt(11)
            else:
                # Fallback al metodo semplice
                text = page_data['text']
                paragraphs = text.split('\n\n')
                
                for para_text in paragraphs:
                    if para_text.strip():
                        p = doc.add_paragraph(para_text)
                        p.style = 'Normal'
                        for run in p.runs:
                            run.font.size = Pt(11)
            
            # Page break (tranne ultima)
            if i < total_pages - 1:
                doc.add_page_break()
        
        if progress_callback:
            progress_callback(90, "Salvataggio DOCX con struttura...")
        
        doc.save(output_path)
        
        logger.info(f"DOCX generato (con preservazione struttura): {output_path}")

